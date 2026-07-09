"""
Parser for American Express "Transaction Details" Excel exports.

This is a different Amex export than AmexParser handles: it has no
'Card Member' / 'Account #' columns. Instead the masked account number
sits in a metadata block above the header row, on a 'Transaction Details'
sheet (alongside a separate 'Transaction Summary' sheet).
"""
import re
import json
import pandas as pd
import openpyxl
from pathlib import Path
from typing import List

from .base_parser import BaseParser
from ..models import Transaction


class AmexActivityParser(BaseParser):
    """Parser for American Express 'Transaction Details' Excel exports"""

    SHEET_NAME = 'Transaction Details'
    HEADER_ROW = 6  # rows to skip to reach the real header row

    EXPECTED_COLUMNS = [
        'Date', 'Description', 'Amount', 'Extended Details',
        'Appears On Your Statement As', 'Reference', 'Category'
    ]

    def __init__(self):
        super().__init__(institution='American Express', account_type='credit_card')

    def _get_sheet_name(self, file_path: Path) -> str:
        try:
            xl = pd.ExcelFile(file_path)
            if self.SHEET_NAME in xl.sheet_names:
                return self.SHEET_NAME
            return xl.sheet_names[0]
        except Exception:
            return self.SHEET_NAME

    def get_account_name(self, file_path: Path) -> str:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            sheet_name = self.SHEET_NAME if self.SHEET_NAME in wb.sheetnames else wb.sheetnames[0]
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(min_row=1, max_row=self.HEADER_ROW + 1, values_only=True))
            for i, row in enumerate(rows):
                if row and row[0] and 'account number' in str(row[0]).lower():
                    if i + 1 < len(rows) and rows[i + 1][0]:
                        digits = re.sub(r'\D', '', str(rows[i + 1][0]))
                        if len(digits) >= 5:
                            return f"Amex *{digits[-5:]}"
                        elif digits:
                            return f"Amex *{digits}"
                    break
        except Exception:
            pass
        return "American Express"

    def detect(self, file_path: Path) -> bool:
        """Detect if this is an Amex 'Transaction Details' Excel export"""
        try:
            if file_path.suffix.lower() not in ['.xlsx', '.xls']:
                return False
            sheet_name = self._get_sheet_name(file_path)
            df = pd.read_excel(file_path, sheet_name=sheet_name, skiprows=self.HEADER_ROW, nrows=0)
            columns = df.columns.tolist()
            return all(col in columns for col in self.EXPECTED_COLUMNS)
        except Exception:
            return False

    def parse(self, file_path: Path, account_name: str) -> List[Transaction]:
        """Parse Amex 'Transaction Details' Excel export"""
        sheet_name = self._get_sheet_name(file_path)
        df = pd.read_excel(file_path, sheet_name=sheet_name, skiprows=self.HEADER_ROW)
        transactions = []

        for _, row in df.iterrows():
            if row.isna().all():
                continue
            if pd.isna(row['Date']):
                continue

            try:
                date = self._parse_date(str(row['Date']))
            except Exception:
                continue  # Skip rows with invalid dates

            # Amex: positive = charge (expense), negative = credit/refund. Flip sign.
            amount = -self._parse_amount(row['Amount'])

            category = row.get('Category', 'Uncategorized')
            if pd.isna(category):
                category = 'Uncategorized'
            # Amex categories come as "Parent-Child" (e.g. "Travel-Lodging").
            # Keep only the top-level parent to avoid over-granular categories.
            category = str(category).split('-')[0].strip()

            reference = row.get('Reference')

            transaction = Transaction(
                date=date,
                description=row['Description'],
                amount=amount,
                account_name=account_name,
                account_type=self.account_type,
                institution=self.institution,
                category=str(category),
                raw_data=json.dumps(row.to_dict(), default=str),
                notes=f"Ref: {reference}" if pd.notna(reference) else None
            )

            transactions.append(transaction)

        return transactions
